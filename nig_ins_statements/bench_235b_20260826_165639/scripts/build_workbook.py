#!/usr/bin/env python3
"""Deterministically build the audited Excel workbook + per-table CSVs from validated JSON.

Worksheets: Index, Page Coverage, Table Inventory, one sheet per table, Review Log, Run Metadata.
Never asks the model for CSV/Excel. Preserves raw + normalized; recalculated totals never overwrite
printed values (checks only). Per workbook-and-qa.md.
"""
import csv
import json
import re
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

RUN_DIR = Path(__file__).resolve().parent.parent
DOC = "2023 sterling FCR FINAL.pdf"
RUN_TS = (RUN_DIR / "RUN_TS.txt").read_text().strip()
OUT_DIR = RUN_DIR / "output"
CSV_DIR = OUT_DIR / "csv"
CSV_DIR.mkdir(parents=True, exist_ok=True)

manifest = json.loads((RUN_DIR / "round1" / "page_manifest.json").read_text())
inventory = json.loads((RUN_DIR / "round1" / "table_inventory.json").read_text())
audit = json.loads((RUN_DIR / "round3" / "document_audit.json").read_text())
validated = {p.stem: json.loads(p.read_text())
             for p in sorted((RUN_DIR / "round3" / "validated_tables").glob("*.json"))}
r2tables = {p.stem: json.loads(p.read_text())
            for p in sorted((RUN_DIR / "round2" / "tables").glob("*.json"))}

AMBER = PatternFill("solid", fgColor="FFE699")
RED = PatternFill("solid", fgColor="FFC7CE")
HDR = PatternFill("solid", fgColor="D9E1F2")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=13)
CITE = Font(italic=True, color="555555")

wb = openpyxl.Workbook()


def ws_name(tid, title):
    base = re.sub(r"[\[\]:*?/\\]", " ", (title or "").strip())
    name = f"{tid} {base}"[:31].strip()
    return name or tid


def cite_line(pages, printed):
    pp = ", ".join(str(x) for x in pages) if pages else "?"
    printed = [x for x in (printed or []) if x]
    ptxt = f"; printed page {', '.join(map(str, printed))}" if printed else ""
    return f"Source: PDF page {pp}{ptxt}"


# map table_id -> printed pages from inventory
inv_by_id = {t["table_id"]: t for t in inventory}
review_rows = []  # collected issues


def add_issue(issue_id, tid, title, sheet, loc, pdf_pages, category, obs, treatment,
              conf, mat, action, status, r3id):
    review_rows.append([issue_id, f"{tid} {title or ''}".strip(), f"{sheet}!{loc}",
                        ", ".join(map(str, pdf_pages)), "", category, obs, treatment,
                        conf, mat, action, status, r3id])


# ---------------- Data worksheets ----------------
manifest_rows = []  # for index
issue_seq = 0
for tid in sorted(validated) or sorted(r2tables):
    v = validated.get(tid) or r2tables[tid]
    is_val = tid in validated
    title = v.get("final_title_raw") if is_val else v.get("title_raw")
    headers = v.get("final_headers_raw") if is_val else v.get("headers_raw")
    rows = v.get("final_normalized_rows") if is_val else v.get("normalized_rows")
    raw_rows = v.get("final_rows_raw") if is_val else v.get("rows_raw")
    verdict = v.get("verdict", "unvalidated")
    src_pages = (v.get("source_pages") if is_val else [sp["pdf_page"] for sp in v.get("source_pages", [])]) or []
    if src_pages and isinstance(src_pages[0], dict):
        src_pages = [sp["pdf_page"] for sp in src_pages]
    printed = inv_by_id.get(tid, {}).get("printed_pages", [])
    units = r2tables.get(tid, {}).get("units_raw")

    sheet = ws_name(tid, title)
    ws = wb.create_sheet(sheet)
    ws["A1"] = title or tid
    ws["A1"].font = TITLE
    ws["A2"] = cite_line(src_pages, printed)
    ws["A2"].font = CITE
    ws["A3"] = f"Units: {units or 'as printed'}   |   Validation verdict: {verdict}"
    ws["A3"].font = Font(italic=True, size=10, color="777777")

    r0 = 5
    if headers:
        for c, h in enumerate(headers, 1):
            cell = ws.cell(r0, c, h)
            cell.font = BOLD
            cell.fill = HDR
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        r0 += 1
    for row in (rows or []):
        for c, val in enumerate(row, 1):
            ws.cell(r0, c, val)
        r0 += 1

    # flag sheet colour by verdict / issues
    flagged = verdict in ("needs_human_review", "rejected_incomplete")
    if flagged:
        ws["A1"].fill = AMBER

    # write CSV (raw printed rows for auditability)
    with open(CSV_DIR / f"{tid}.csv", "w", newline="") as f:
        w = csv.writer(f)
        w.writerow([f"# {title or tid} | {cite_line(src_pages, printed)} | verdict={verdict}"])
        if headers:
            w.writerow(headers)
        for row in (raw_rows or []):
            w.writerow(row)

    # column widths
    for c in range(1, (len(headers) if headers else 6) + 1):
        ws.column_dimensions[get_column_letter(c)].width = 20

    # collect issues into review log
    if is_val:
        for corr in v.get("corrections", []):
            issue_seq += 1
            add_issue(f"C{issue_seq:03d}", tid, title, sheet,
                      f"r{corr.get('row_index')}c{corr.get('column_index')}",
                      [corr.get("source_pdf_page")], "extraction_correction",
                      f"R2='{corr.get('round2_value')}' -> validated='{corr.get('validated_value')}': {corr.get('reason')}",
                      "Corrected in workbook", corr.get("confidence", ""), "review_recommended",
                      "Confirm correction against source", "Resolved/Preserved as printed", f"r3v-{tid}")
        for iss in v.get("issues", []):
            issue_seq += 1
            add_issue(f"I{issue_seq:03d}", tid, title, sheet, iss.get("location", ""),
                      src_pages, iss.get("category", ""), iss.get("observation", ""),
                      "Flagged", iss.get("confidence", ""), iss.get("materiality", ""),
                      iss.get("recommended_action", ""),
                      "Unresolved" if iss.get("materiality") == "material" else "Preserved as printed",
                      f"r3v-{tid}")
        checks = v.get("checks", {})
        if checks.get("totals_status") == "does_not_reconcile":
            issue_seq += 1
            add_issue(f"A{issue_seq:03d}", tid, title, sheet, "totals", src_pages,
                      "arithmetic_anomaly", "Printed totals/ratios do not reconcile with components",
                      "Preserved as printed (not overwritten)", "high", "review_recommended",
                      "Check printed totals against components", "Preserved as printed", f"r3v-{tid}")

    manifest_rows.append([tid, title or "", sheet, cite_line(src_pages, printed), verdict])

# document-audit issues into review log
for iss in audit.get("issues", []):
    issue_seq += 1
    add_issue(iss.get("issue_id", f"D{issue_seq:03d}"), ",".join(iss.get("table_ids", [])), "",
              "Table Inventory", "", iss.get("pdf_pages", []), iss.get("category", ""),
              f"{iss.get('observation','')} | evidence: {iss.get('evidence','')}",
              "Audit finding", iss.get("confidence", ""), iss.get("materiality", ""),
              iss.get("recommended_action", ""),
              {"resolved_from_document": "Resolved", "preserved_as_printed": "Preserved as printed",
               "unresolved": "Unresolved"}.get(iss.get("status"), "Unresolved"),
              "r3-audit")

# ---------------- Index ----------------
idx = wb["Sheet"]
idx.title = "Index"
idx["A1"] = "Scanned-Table Extraction Workbook"
idx["A1"].font = TITLE
meta = [
    ("Source file", DOC),
    ("Run ID", RUN_TS),
    ("Extraction date", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ("Model (all rounds)", "Qwen3-VL-235B-A22B-Instruct-FP8"),
    ("PDF pages", manifest["expected_pages"]),
    ("Logical tables", len(inventory)),
    ("Delivered tables", len(manifest_rows)),
    ("Notation", "Raw printed values kept in per-table CSVs; normalized values in worksheets."),
    ("Grouping", "One worksheet per logical table; repeated versions kept separate."),
]
for i, (k, val) in enumerate(meta, 3):
    idx.cell(i, 1, k).font = BOLD
    idx.cell(i, 2, val)
hdr_r = 3 + len(meta) + 1
for c, h in enumerate(["Table ID", "Title", "Worksheet", "Citation", "Verdict"], 1):
    cell = idx.cell(hdr_r, c, h); cell.font = BOLD; cell.fill = HDR
for i, row in enumerate(manifest_rows, hdr_r + 1):
    for c, val in enumerate(row, 1):
        idx.cell(i, c, val)
for c, wdt in zip(range(1, 6), (12, 45, 26, 40, 18)):
    idx.column_dimensions[get_column_letter(c)].width = wdt

# ---------------- Page Coverage ----------------
pc = wb.create_sheet("Page Coverage")
for c, h in enumerate(["PDF page", "Printed page", "Round1 status", "Legibility",
                       "Table IDs", "R3 disposition", "Exclusion reason"], 1):
    cell = pc.cell(1, c, h); cell.font = BOLD; cell.fill = HDR
tables_by_page = {}
for t in inventory:
    for pg in t["source_pages"]:
        tables_by_page.setdefault(pg, []).append(t["table_id"])
verdict_by_tid = {tid: v.get("verdict") for tid, v in validated.items()}
for i, p in enumerate(manifest["pages"], 2):
    pg = p["pdf_page"]
    tids = tables_by_page.get(pg, [])
    disp = ", ".join(f"{t}:{verdict_by_tid.get(t,'-')}" for t in tids) if tids else ""
    excl = "no table on page" if p["page_status"] == "no_table" else ""
    for c, val in enumerate([pg, p.get("printed_page"), p["page_status"], p["page_legibility"],
                             ", ".join(tids), disp, excl], 1):
        pc.cell(i, c, val)
for c, wdt in zip(range(1, 8), (9, 12, 16, 11, 16, 34, 18)):
    pc.column_dimensions[get_column_letter(c)].width = wdt

# ---------------- Table Inventory ----------------
ti = wb.create_sheet("Table Inventory")
for c, h in enumerate(["Table ID", "Title", "Units", "Source pages", "Repeat of",
                       "Complexity", "Worksheet", "Verdict"], 1):
    cell = ti.cell(1, c, h); cell.font = BOLD; cell.fill = HDR
row_lookup = {m[0]: m for m in manifest_rows}
for i, t in enumerate(inventory, 2):
    m = row_lookup.get(t["table_id"], [t["table_id"], t["title_raw"], "", "", "-"])
    for c, val in enumerate([t["table_id"], t["title_raw"], t.get("units_raw"),
                             ", ".join(map(str, t["source_pages"])), t.get("possible_repeat_of") or "",
                             t.get("complexity"), m[2], m[4]], 1):
        ti.cell(i, c, val)
for c, wdt in zip(range(1, 9), (10, 42, 14, 14, 12, 12, 24, 18)):
    ti.column_dimensions[get_column_letter(c)].width = wdt

# ---------------- Review Log ----------------
rl = wb.create_sheet("Review Log")
cols = ["Issue ID", "Table", "Sheet!cell", "PDF pages", "Printed pages", "Category",
        "Source observation", "Workbook treatment", "Confidence", "Materiality",
        "User action/check", "Status", "R3 request ID"]
for c, h in enumerate(cols, 1):
    cell = rl.cell(1, c, h); cell.font = BOLD; cell.fill = HDR
for i, r in enumerate(review_rows, 2):
    for c, val in enumerate(r, 1):
        cell = rl.cell(i, c, val)
        if r[9] == "material":
            cell.fill = RED
        elif r[9] == "review_recommended":
            cell.fill = AMBER
for c, wdt in zip(range(1, 14), (9, 26, 16, 11, 12, 20, 46, 22, 11, 16, 30, 20, 14)):
    rl.column_dimensions[get_column_letter(c)].width = wdt

# ---------------- Run Metadata ----------------
rm = wb.create_sheet("Run Metadata")
usage = {"round1": 0, "round2": 0, "round3": 0}
prompt_tok = {"round1": 0, "round2": 0, "round3": 0}
for rnd in usage:
    raw = RUN_DIR / rnd / f"{rnd}_raw.jsonl"
    if raw.exists():
        for line in raw.read_text().splitlines():
            if not line.strip():
                continue
            u = (json.loads(line).get("response") or {}).get("body", {}).get("usage") or {}
            usage[rnd] += u.get("completion_tokens", 0)
            prompt_tok[rnd] += u.get("prompt_tokens", 0)
img_manifest = json.loads((RUN_DIR / "round1" / "image_manifest.json").read_text())
mrows = [
    ("Run ID", RUN_TS),
    ("Built", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ("R1 model / schema", "Qwen3-VL-235B-A22B-Instruct-FP8 / round1-discovery"),
    ("R2 model / schema", "Qwen3-VL-235B-A22B-Instruct-FP8 / round2-extraction"),
    ("R3 model / schema", "Qwen3-VL-235B-A22B-Instruct-FP8 / round3-table-validation + round3-document-audit"),
    ("Pages imaged (discovery 150dpi jpg / extraction 200dpi q88 jpg)", len(img_manifest)),
    ("R1 prompt/completion tokens", f"{prompt_tok['round1']} / {usage['round1']}"),
    ("R2 prompt/completion tokens", f"{prompt_tok['round2']} / {usage['round2']}"),
    ("R3 prompt/completion tokens", f"{prompt_tok['round3']} / {usage['round3']}"),
    ("Image hashes", "see round1/image_manifest.json (sha256 per page)"),
    ("Raw results preserved", "round{1,2,3}/*_raw.jsonl"),
]
for i, (k, val) in enumerate(mrows, 1):
    rm.cell(i, 1, k).font = BOLD
    rm.cell(i, 2, val)
rm.column_dimensions["A"].width = 52
rm.column_dimensions["B"].width = 60

# order sheets: Index, Page Coverage, Table Inventory, data..., Review Log, Run Metadata
order = ["Index", "Page Coverage", "Table Inventory"]
data_sheets = [s for s in wb.sheetnames if s not in order + ["Review Log", "Run Metadata"]]
wb._sheets.sort(key=lambda s: (order + data_sheets + ["Review Log", "Run Metadata"]).index(s.title))

out_xlsx = OUT_DIR / f"sterling_fcr_2023_tables_{RUN_TS}.xlsx"
wb.save(out_xlsx)
print(f"Workbook: {out_xlsx}")
print(f"Worksheets: {len(wb.sheetnames)} | tables: {len(manifest_rows)} | review items: {len(review_rows)} | CSVs: {len(list(CSV_DIR.glob('*.csv')))}")
